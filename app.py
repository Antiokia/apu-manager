import streamlit as st
import openpyxl
import re
import io
import os
from collections import defaultdict

# --- IMPORTAMOS O REPLICAMOS TUS HELPERS ---
def rgb_to_hex(r, g, b): return f"{r:02X}{g:02X}{b:02X}"
def safe_float(val):
    try: return float(val) if val not in (None, "", " ") else 0.0
    except: return 0.0

def thin_border():
    s = openpyxl.styles.Side(style="thin")
    return openpyxl.styles.Border(left=s, right=s, top=s, bottom=s)

def apply_header_style(cell, hex_color="00467F", font_color="FFFFFF"):
    cell.font = openpyxl.styles.Font(bold=True, color=font_color, name="Arial")
    cell.fill = openpyxl.styles.PatternFill("solid", fgColor=hex_color)
    cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")

def money_fmt(cell): cell.number_format = '$ #,##0'
def qty_fmt(cell): cell.number_format = '#,##0.00'
def pct_fmt(cell): cell.number_format = '0.00%'

# --- TUS FUNCIONES DE LÓGICA (Adaptadas para memoria) ---
def leer_categorias(ws_aux):
    cats = []
    for row in ws_aux.iter_rows(min_row=2, values_only=True):
        texto_apu = str(row[0] or "").strip()
        if not texto_apu: continue
        categoria = str(row[1] or "").strip()
        num       = str(row[2] or "").strip()
        color_str = str(row[3] or "220,220,220").strip()
        tipo      = str(row[4] or "ESTANDAR").strip().upper()
        parts = [p.strip() for p in color_str.split(",")]
        try: r, g, b = int(parts[0]), int(parts[1]), int(parts[2])
        except: r, g, b = 220, 220, 220
        cats.append({"texto_apu": texto_apu, "categoria": categoria, "num": num, "hex": rgb_to_hex(r, g, b), "tipo": tipo})
    return cats

def leer_cant_ppto(ws_ppto):
    cant_map = {}
    for row in ws_ppto.iter_rows(min_row=3, max_row=200, values_only=True):
        item_val = row[0]
        if item_val is None: continue
        try:
            item_num = int(item_val)
            cant_map[item_num] = safe_float(row[3])
        except: pass
    return cant_map

def detectar_categoria(cell_a, cats):
    for cat in cats:
        if cat["texto_apu"].upper() in cell_a.upper(): return cat["categoria"], cat["tipo"]
    return None, None

def procesar_item(ws_item, item_num, cant_ppto, cats, diccionarios):
    if cant_ppto == 0: return 0.0
    categoria_actual, tipo_actual, base_item = None, None, 0.0
    
    d_insumos, d_unidades, d_cantidades, d_valores, d_categorias, d_vunit, d_claves = (
        diccionarios["insumos"], diccionarios["unidades"], diccionarios["cantidades"],
        diccionarios["valores"], diccionarios["categorias"], diccionarios["vunit"], diccionarios["claves"]
    )

    for row in ws_item.iter_rows(min_row=11, max_row=ws_item.max_row, values_only=True):
        cell_a = str(row[0] or "").strip()
        cell_b = str(row[1] or "").strip()
        cell_f = str(row[5] or "").strip()

        cat, tipo = detectar_categoria(cell_a, cats)
        if cat: categoria_actual, tipo_actual = cat, tipo
        if "TOTAL COSTO" in cell_a.upper(): categoria_actual, tipo_actual = None, None; continue
        if categoria_actual is None or not cell_b or "DESCRIPCI" in cell_b.upper() or "CLAVE" in cell_a.upper() or "SUB-TOTAL" in cell_a.upper() or "SUB-TOTAL" in cell_f.upper(): continue

        clave_insumo, descripcion, unidad = cell_a, cell_b, str(row[2] or "").strip() or "UND"
        cant_insumo, v_unit = 0.0, 0.0

        if tipo_actual == "RENDIMIENTO":
            cant_insumo = safe_float(row[3]) * (safe_float(row[4]) or 1.0)
            v_unit = safe_float(row[5])
        elif tipo_actual == "PORCENTAJE":
            cant_insumo, v_unit, unidad = safe_float(row[4]), safe_float(row[5]), str(row[2] or "").strip() or "GL"
        else:
            cant_insumo, v_unit = safe_float(row[3]), safe_float(row[5])

        if v_unit <= 0: continue
        clave = descripcion.upper().strip()
        valor_total = cant_insumo * v_unit

        if tipo_actual != "PORCENTAJE": base_item += valor_total * cant_ppto

        if clave in d_insumos:
            d_cantidades[clave] += cant_insumo * cant_ppto
            d_valores[clave]    += valor_total * cant_ppto
            if d_cantidades[clave] > 0: d_vunit[clave] = d_valores[clave] / d_cantidades[clave]
            if not d_claves[clave] and clave_insumo: d_claves[clave] = clave_insumo
        else:
            d_insumos[clave], d_unidades[clave], d_cantidades[clave], d_valores[clave], d_categorias[clave], d_vunit[clave], d_claves[clave] = descripcion, unidad, cant_insumo * cant_ppto, valor_total * cant_ppto, categoria_actual, v_unit, clave_insumo
    return base_item

def escribir_explosion(wb, cats, diccionarios, base_total):
    sheet_name = "EXPLOSION INSUMOS"
    if sheet_name in wb.sheetnames: del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    
    d_insumos, d_unidades, d_cantidades, d_valores, d_categorias, d_vunit, d_claves = (
        diccionarios["insumos"], diccionarios["unidades"], diccionarios["cantidades"],
        diccionarios["valores"], diccionarios["categorias"], diccionarios["vunit"], diccionarios["claves"]
    )

    ws.merge_cells("A1:G1")
    ws["A1"] = "EXPLOSIÓN DE INSUMOS - CONSOLIDADO"
    ws["A1"].font, ws["A1"].alignment = openpyxl.styles.Font(bold=True, size=14, name="Arial"), openpyxl.styles.Alignment(horizontal="center")

    headers = ["CATEGORÍA", "CLAVE", "DESCRIPCIÓN DEL INSUMO", "UNIDAD", "CANTIDAD TOTAL", "COSTO", "VALOR TOTAL"]
    for col_idx, h in enumerate(headers, start=1):
        apply_header_style(ws.cell(row=3, column=col_idx, value=h))
        ws.cell(row=3, column=col_idx).border = thin_border()

    col_widths = [18, 12, 55, 10, 18, 18, 18]
    for i, w in enumerate(col_widths, start=1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    fila, total_gral = 4, 0.0
    for cat in cats:
        cat_nombre, cat_hex, cat_tipo, cat_label = cat["categoria"], cat["hex"], cat["tipo"], f"{cat['num']} {cat['categoria']}"
        insumos_cat = [(k, v) for k, v in d_insumos.items() if d_categorias.get(k) == cat_nombre]
        if not insumos_cat: continue

        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=7)
        cell = ws.cell(row=fila, column=1, value=cat_label)
        cell.font, cell.fill = openpyxl.styles.Font(bold=True, name="Arial"), openpyxl.styles.PatternFill("solid", fgColor=cat_hex)
        fila += 1

        subtotal = 0.0
        for clave, descripcion in insumos_cat:
            cant_total, valor_total = d_cantidades[clave], d_valores[clave]
            v_unit_mostrar = base_total if cat_tipo == "PORCENTAJE" else (valor_total / cant_total if cant_total > 0 else d_vunit[clave])
            if cat_tipo == "PORCENTAJE" and base_total > 0: cant_total = valor_total / base_total

            row_data = [cat_nombre, d_claves[clave], descripcion, d_unidades[clave], cant_total, v_unit_mostrar, valor_total]
            for col_idx, val in enumerate(row_data, start=1):
                c = ws.cell(row=fila, column=col_idx, value=val)
                c.border, c.font = thin_border(), openpyxl.styles.Font(name="Arial")

            pct_fmt(ws.cell(row=fila, column=5)) if cat_tipo == "PORCENTAJE" else qty_fmt(ws.cell(row=fila, column=5))
            money_fmt(ws.cell(row=fila, column=6))
            money_fmt(ws.cell(row=fila, column=7))
            subtotal += valor_total; fila += 1

        ws.cell(row=fila, column=6, value=f"SUBTOTAL {cat_nombre}:").font = openpyxl.styles.Font(bold=True, name="Arial")
        money_fmt(ws.cell(row=fila, column=7, value=subtotal))
        ws.cell(row=fila, column=7).font = openpyxl.styles.Font(bold=True, name="Arial")
        for col_idx in range(1, 8):
            ws.cell(row=fila, column=col_idx).fill = openpyxl.styles.PatternFill("solid", fgColor=cat_hex)
            ws.cell(row=fila, column=col_idx).border = thin_border()
        total_gral += subtotal; fila += 2

    ws.cell(row=fila, column=6, value="TOTAL GENERAL:").font = openpyxl.styles.Font(bold=True, color="FFFFFF", name="Arial")
    money_fmt(ws.cell(row=fila, column=7, value=total_gral))
    ws.cell(row=fila, column=7).font = openpyxl.styles.Font(bold=True, color="FFFFFF", name="Arial")
    for col_idx in range(1, 8):
        ws.cell(row=fila, column=col_idx).fill = openpyxl.styles.PatternFill("solid", fgColor="00467F")
        ws.cell(row=fila, column=col_idx).border = thin_border()
    return total_gral


# --- INTERFAZ GRÁFICA CON STREAMLIT ---
st.set_page_config(page_title="Explosión de Insumos", page_icon="📊", layout="centered")

st.title("📊 Generador de Explosión de Insumos")
st.write("Carga tu archivo `.xlsm` de presupuesto con APU para consolidar automáticamente los insumos.")

uploaded_file = st.file_uploader("Elige tu archivo de Excel (.xlsm)", type=["xlsm"])

if uploaded_file is not None:
    st.success("Archivo cargado con éxito en la memoria.")
    
    if st.button("🚀 Procesar Explosión de Insumos"):
        with st.spinner("Procesando datos y estructurando el reporte..."):
            try:
                # Leer el archivo desde la memoria de carga
                wb = openpyxl.load_workbook(uploaded_file, data_only=True)
                
                ws_aux = wb["AUX"] if "AUX" in wb.sheetnames else (wb["Aux"] if "Aux" in wb.sheetnames else None)
                ws_ppto = wb["PPTO"] if "PPTO" in wb.sheetnames else None
                
                if ws_aux is None or ws_ppto is None:
                    st.error("Error: El archivo debe contener obligatoriamente las pestañas 'AUX' y 'PPTO'.")
                else:
                    cats = leer_categorias(ws_aux)
                    cant_ppto_map = leer_cant_ppto(ws_ppto)
                    
                    diccionarios = {
                        "insumos": {}, "unidades": {}, "cantidades": defaultdict(float),
                        "valores": defaultdict(float), "categorias": {}, "vunit": {}, "claves": defaultdict(str)
                    }
                    
                    item_pattern = re.compile(r"^item\s+(\d+)$", re.IGNORECASE)
                    base_total = 0.0
                    
                    for sheet_name in wb.sheetnames:
                        m = item_pattern.match(sheet_name.strip())
                        if m:
                            item_num = int(m.group(1))
                            cant_ppto = cant_ppto_map.get(item_num, 0.0)
                            base_total += procesar_item(wb[sheet_name], item_num, cant_ppto, cats, diccionarios)
                    
                    # Generamos la hoja dentro de nuestro libro abierto
                    total_gral = escribir_explosion(wb, cats, diccionarios, base_total)
                    
                    # Guardamos el resultado en un flujo de bytes en memoria (evita errores de rutas en Mac)
                    output = io.BytesIO()
                    wb.save(output)
                    output.seek(0)
                    
                    st.balloons()
                    st.metric(label="COSTO TOTAL GENERAL CONSOLIDADO", value=f"$ {total_gral:,.2f}")
                    
                    # Botón para que el usuario descargue el archivo final
                    st.download_button(
                        label="📥 Descargar Reporte Generado",
                        data=output,
                        file_name=f"{os.path.splitext(uploaded_file.name)[0]}_explosion.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
            except Exception as e:
                st.error(f"Ocurrió un error inesperado al procesar: {e}")